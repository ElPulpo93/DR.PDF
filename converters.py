"""Conversion engine for Dr.PDF.

Everything runs on this computer. Word / Excel / PowerPoint files are
converted through the copy of Microsoft Office installed here (or
LibreOffice, if that is what's installed). Nothing is uploaded anywhere.
"""
import csv
import os
import shutil
import subprocess
import tempfile
import winreg
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageOps

try:  # HEIC / HEIF (iPhone photos) support
    import pillow_heif
    pillow_heif.register_heif_opener()
    HEIF_OK = True
except Exception:
    HEIF_OK = False


# ---------------------------------------------------------------- detection
def _com_registered(progid):
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, progid + r"\CLSID"))
        return True
    except OSError:
        return False


def _find_soffice():
    for base in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
        if base:
            p = Path(base) / "LibreOffice" / "program" / "soffice.exe"
            if p.exists():
                return str(p)
    return shutil.which("soffice")


HAS_WORD = _com_registered("Word.Application")
HAS_EXCEL = _com_registered("Excel.Application")
HAS_PPT = _com_registered("PowerPoint.Application")
SOFFICE = _find_soffice()


# ---------------------------------------------------------------- formats
# target label -> (extension, description)
TARGETS = {
    "PDF":  (".pdf", "PDF document"),
    "DOCX": (".docx", "Word document"),
    "XLSX": (".xlsx", "Excel workbook"),
    "PPTX": (".pptx", "PowerPoint presentation"),
    "CSV":  (".csv", "Comma-separated values (one file per sheet)"),
    "TXT":  (".txt", "Plain text"),
    "RTF":  (".rtf", "Rich Text Format"),
    "ODT":  (".odt", "OpenDocument text"),
    "ODS":  (".ods", "OpenDocument spreadsheet"),
    "ODP":  (".odp", "OpenDocument presentation"),
    "HTML": (".html", "Web page"),
    "JPG":  (".jpg", "JPEG picture"),
    "PNG":  (".png", "PNG picture"),
    "WEBP": (".webp", "WebP picture"),
    "BMP":  (".bmp", "Bitmap picture"),
    "TIFF": (".tiff", "TIFF picture"),
    "GIF":  (".gif", "GIF picture"),
    "ICO":  (".ico", "Windows icon"),
}
if HEIF_OK:
    TARGETS["HEIC"] = (".heic", "HEIC picture (iPhone)")

# image target label -> (Pillow format, supports quality)
IMAGE_FMT = {
    "JPG": ("JPEG", True), "PNG": ("PNG", False), "WEBP": ("WEBP", True),
    "BMP": ("BMP", False), "TIFF": ("TIFF", False), "GIF": ("GIF", False),
    "ICO": ("ICO", False), "HEIC": ("HEIF", True),
}
QUALITY_TARGETS = {k for k, (_, q) in IMAGE_FMT.items() if q}
IMAGE_TARGETS = [k for k in IMAGE_FMT if k in TARGETS]

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".jfif", ".webp", ".bmp", ".gif",
              ".tif", ".tiff", ".ico", ".tga", ".ppm", ".pgm", ".pbm"}
if HEIF_OK:
    IMAGE_EXTS |= {".heic", ".heif", ".avif"}
WORD_EXTS = {".docx", ".doc", ".docm", ".dotx", ".rtf", ".odt", ".htm", ".html"}
TEXT_EXTS = {".txt", ".md", ".log"}
SHEET_EXTS = {".xlsx", ".xls", ".xlsm", ".xlsb", ".ods"}
CSV_EXTS = {".csv", ".tsv"}
SLIDE_EXTS = {".pptx", ".ppt", ".pptm", ".ppsx", ".pps", ".odp"}

# extension that a target label produces, used to hide "docx -> DOCX" etc.
_SAME = {".jpeg": "JPG", ".jpg": "JPG", ".docx": "DOCX", ".doc": "DOC",
         ".rtf": "RTF", ".odt": "ODT", ".htm": "HTML", ".html": "HTML",
         ".xlsx": "XLSX", ".xls": "XLS", ".ods": "ODS", ".csv": "CSV",
         ".pptx": "PPTX", ".ppt": "PPT", ".odp": "ODP", ".txt": "TXT"}

_OFFICE_DOC = ["PDF", "DOCX", "RTF", "ODT", "TXT", "HTML"]
_OFFICE_SHEET = ["PDF", "XLSX", "CSV", "ODS"]
_OFFICE_SLIDES = ["PDF", "PPTX", "ODP", "PNG", "JPG"]


def kind_of(ext):
    ext = ext.lower()
    if ext in IMAGE_EXTS:
        return "image"
    if ext == ".pdf":
        return "pdf"
    if ext in WORD_EXTS:
        return "word"
    if ext in TEXT_EXTS:
        return "text"
    if ext in SHEET_EXTS:
        return "sheet"
    if ext in CSV_EXTS:
        return "csv"
    if ext in SLIDE_EXTS:
        return "slides"
    return None


def targets_for(ext):
    """Ordered list of target labels a file with this extension can become."""
    ext = ext.lower()
    kind = kind_of(ext)
    office = bool(SOFFICE)
    if kind == "image":
        t = IMAGE_TARGETS + ["PDF", "DOCX"]
    elif kind == "pdf":
        t = ["DOCX", "XLSX", "TXT", "JPG", "PNG", "PDF"]
    elif kind == "word":
        if HAS_WORD or office:
            t = list(_OFFICE_DOC)
        elif ext == ".docx":
            t = ["TXT"]
        else:
            t = []
    elif kind == "text":
        t = ["PDF", "DOCX"]
    elif kind == "sheet":
        if HAS_EXCEL or office:
            t = list(_OFFICE_SHEET)
        elif ext in (".xlsx", ".xlsm"):
            t = ["CSV", "PDF"]
        else:
            t = []
    elif kind == "csv":
        t = ["XLSX", "PDF"] + (["CSV"] if ext == ".tsv" else [])
    elif kind == "slides":
        t = list(_OFFICE_SLIDES) if (HAS_PPT or office) else []
    else:
        t = []
    if kind != "image":
        same = _SAME.get(ext)
        t = [x for x in t if x != same or x == "PDF"]
    return [x for x in t if x in TARGETS]


def is_supported(path: Path):
    return bool(targets_for(path.suffix))


def needs_office_hint(path: Path):
    """True when a file is an Office type we can't handle without Office."""
    k = kind_of(path.suffix)
    return k in ("word", "sheet", "slides") and not targets_for(path.suffix)


# ---------------------------------------------------------------- helpers
def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    i = 1
    while True:
        candidate = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not candidate.exists():
            return candidate
        i += 1


def _num(s):
    """Turn spreadsheet-looking text into a number where it is safe to."""
    if s is None:
        return None
    t = str(s).strip()
    if not t:
        return None
    if t.startswith("0") and len(t) > 1 and not t.startswith("0."):
        return t  # keep zip codes / IDs like 00123
    c = t.replace(",", "").replace("$", "")
    neg = c.startswith("(") and c.endswith(")")
    if neg:
        c = c[1:-1]
    pct = c.endswith("%")
    if pct:
        c = c[:-1]
    try:
        v = float(c)
    except ValueError:
        return t
    if neg:
        v = -v
    if pct:
        return v / 100
    return int(v) if v.is_integer() and "." not in c and "e" not in c.lower() else v


def _autofit(ws):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                n = max(len(x) for x in str(cell.value).split("\n"))
                widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), n)
    for col, w in widths.items():
        ws.column_dimensions[col].width = min(max(8, w + 2), 60)


def _font():
    """A Unicode-capable font for text -> PDF, registered once."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "DrBody" in pdfmetrics.getRegisteredFontNames():
        return "DrBody"
    for name in ("segoeui.ttf", "arial.ttf", "calibri.ttf"):
        p = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / name
        if p.exists():
            try:
                pdfmetrics.registerFont(TTFont("DrBody", str(p)))
                return "DrBody"
            except Exception:
                pass
    return "Helvetica"


# ---------------------------------------------------------------- images
def _prepare_image(im, pil_fmt):
    has_alpha = im.mode in ("RGBA", "LA") or (
        im.mode == "P" and "transparency" in im.info)
    if pil_fmt in ("JPEG", "PDF") or (pil_fmt == "BMP" and has_alpha):
        if has_alpha:
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, (255, 255, 255))
            bg.paste(im, mask=im.getchannel("A"))
            im = bg
        elif im.mode != "RGB":
            im = im.convert("RGB")
    elif pil_fmt == "GIF":
        im = im.convert("RGBA" if has_alpha else "RGB").convert(
            "P", palette=Image.ADAPTIVE)
    elif im.mode not in ("RGB", "RGBA", "L", "LA", "P", "1"):
        im = im.convert("RGBA" if has_alpha else "RGB")
    elif im.mode == "P" and pil_fmt in ("WEBP", "HEIF", "ICO"):
        im = im.convert("RGBA" if has_alpha else "RGB")
    return im


def _save_image(im, out: Path, label, quality, exif=None):
    pil_fmt = IMAGE_FMT[label][0] if label in IMAGE_FMT else "PDF"
    im = _prepare_image(im, pil_fmt)
    kwargs = {}
    if pil_fmt in ("JPEG", "WEBP", "HEIF"):
        kwargs["quality"] = quality
    if pil_fmt in ("JPEG", "PNG"):
        kwargs["optimize"] = True
    if pil_fmt == "PDF":
        kwargs["resolution"] = 150
    if pil_fmt == "ICO":
        sizes = [s for s in (16, 24, 32, 48, 64, 128, 256) if s <= max(im.size)] or [16]
        kwargs["sizes"] = [(s, s) for s in sizes]
    if exif and pil_fmt in ("JPEG", "WEBP", "PNG", "TIFF"):
        kwargs["exif"] = exif
    im.save(out, pil_fmt, **kwargs)


def image_convert(src, label, out_dir, opts):
    ext = TARGETS[label][0]
    out = unique_path(out_dir / (src.stem + ext))
    with Image.open(src) as im:
        im.load()
        exif = None if opts.get("strip_meta", True) else im.info.get("exif")
        im = ImageOps.exif_transpose(im)  # keep photos the right way up
        if label == "DOCX":
            _image_to_docx(im, out)
        else:
            _save_image(im, out, label, opts.get("quality", 90), exif)
    return [out]


def _image_to_docx(im, out):
    from docx import Document
    from docx.shared import Inches
    tmp = Path(tempfile.mkdtemp()) / "img.png"
    try:
        _prepare_image(im, "PNG").save(tmp)
        doc = Document()
        sec = doc.sections[0]
        usable = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
        w_in = min(usable, im.width / 96)
        doc.add_picture(str(tmp), width=Inches(w_in))
        doc.save(out)
    finally:
        shutil.rmtree(tmp.parent, ignore_errors=True)


# ---------------------------------------------------------------- PDF input
def pdf_to_docx(src, out):
    from pdf2docx import Converter
    cv = Converter(str(src))
    try:
        cv.convert(str(out))
    finally:
        cv.close()


def _text_grid(page):
    """Rebuild rows/columns from word positions on a page without table borders."""
    words = page.extract_words(keep_blank_chars=False, use_text_flow=True)
    lines = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) < 3:
            lines[-1].append(w)
        else:
            lines.append([w])
    rows = []  # each row: list of (x0, x1, text) chunks
    for ln in lines:
        ln.sort(key=lambda w: w["x0"])
        chunks = []
        for w in ln:
            gap = 0.6 * (w["bottom"] - w["top"])  # wider than a space = new cell
            if chunks and w["x0"] - chunks[-1][1] <= gap:
                x0, _, t = chunks[-1]
                chunks[-1] = (x0, w["x1"], t + " " + w["text"])
            else:
                chunks.append((w["x0"], w["x1"], w["text"]))
        rows.append(chunks)
    if not rows:
        return []
    width = max(len(r) for r in rows)
    full = [r for r in rows if len(r) == width]
    centers = [sum((c[0] + c[1]) / 2 for c in col) / len(col) for col in zip(*full)]
    grid = []
    for r in rows:
        cells = [None] * width
        if len(r) == width:
            cells = [c[2] for c in r]
        else:  # partial row: drop each chunk into the nearest free column
            for x0, x1, t in r:
                order = sorted(range(width), key=lambda k: abs(centers[k] - (x0 + x1) / 2))
                k = next((k for k in order if cells[k] is None), None)
                if k is not None:
                    cells[k] = t
        grid.append(cells)
    return [grid]


def pdf_to_xlsx(src, out):
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    found_any = False
    with pdfplumber.open(str(src)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables()
            if not tables:  # no ruled tables: line up text by its columns
                tables = _text_grid(page)
            tables = [t for t in tables if any(any(c for c in r) for r in t)]
            if not tables:
                continue
            found_any = True
            ws = wb.create_sheet(f"Page {i}")
            r = 1
            for t in tables:
                for j, row in enumerate(t):
                    for c, val in enumerate(row, 1):
                        cell = ws.cell(row=r, column=c, value=_num(val))
                        if j == 0:
                            cell.font = Font(bold=True)
                    r += 1
                r += 1  # blank row between tables
            _autofit(ws)
    if not found_any:
        raise RuntimeError("No text or tables found (is this a scanned PDF?)")
    wb.save(out)


def pdf_to_text(src, out):
    import pymupdf as fitz
    with fitz.open(str(src)) as doc:
        text = "\n\n".join(p.get_text() for p in doc)
    if not text.strip():
        raise RuntimeError("No text found (is this a scanned PDF?)")
    out.write_text(text, encoding="utf-8")


def pdf_to_images(src, label, out_dir, opts):
    import pymupdf as fitz
    outs = []
    with fitz.open(str(src)) as doc:
        n = doc.page_count
        if n > 1:
            out_dir = unique_path(out_dir / f"{src.stem} - pages")
            out_dir.mkdir(parents=True)
        for i, page in enumerate(doc, 1):
            pix = page.get_pixmap(dpi=200, alpha=False)
            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            name = f"{src.stem} - page {i}" if n > 1 else src.stem
            out = unique_path(out_dir / (name + TARGETS[label][0]))
            _save_image(im, out, label, opts.get("quality", 90))
            outs.append(out)
    return outs


# ---------------------------------------------------------------- text input
def _read_text(src):
    data = src.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("latin-1")


def text_to_pdf(src, out):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    style = ParagraphStyle("body", fontName=_font(), fontSize=10.5, leading=14)
    story = []
    for line in _read_text(src).splitlines():
        if line.strip():
            lead = len(line) - len(line.lstrip(" "))
            story.append(Paragraph("&nbsp;" * lead + escape(line.strip()), style))
        else:
            story.append(Spacer(1, 10))
    SimpleDocTemplate(str(out), pagesize=letter, leftMargin=inch, rightMargin=inch,
                      topMargin=inch, bottomMargin=inch,
                      title=src.stem).build(story or [Spacer(1, 1)])


def text_to_docx(src, out):
    from docx import Document
    doc = Document()
    for line in _read_text(src).splitlines():
        doc.add_paragraph(line)
    doc.save(out)


def docx_to_text(src, out):
    from docx import Document
    doc = Document(str(src))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append("\t".join(c.text for c in row.cells))
    out.write_text("\n".join(parts), encoding="utf-8")


# ---------------------------------------------------------------- CSV / sheets
def _read_csv_rows(src):
    text = _read_text(src)
    if src.suffix.lower() == ".tsv":
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(text[:20000], ",;\t|").delimiter
        except csv.Error:
            delim = ","
    return list(csv.reader(text.splitlines(), delimiter=delim))


def csv_to_xlsx(src, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = src.stem[:31] or "Sheet1"
    for i, row in enumerate(_read_csv_rows(src)):
        ws.append([_num(v) for v in row])
        if i == 0:
            for c in ws[1]:
                c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    _autofit(ws)
    wb.save(out)


def csv_to_csv(src, out):
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(_read_csv_rows(src))


def rows_to_pdf(rows, out, title):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    font = _font()
    style = ParagraphStyle("cell", fontName=font, fontSize=8, leading=10)
    width = max((len(r) for r in rows), default=1)
    data = [[Paragraph(escape("" if v is None else str(v)), style) for v in r]
            + [""] * (width - len(r)) for r in rows] or [[""]]
    usable = landscape(letter)[0] - inch
    t = Table(data, colWidths=[usable / width] * width, repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EAF6")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    SimpleDocTemplate(str(out), pagesize=landscape(letter), leftMargin=inch / 2,
                      rightMargin=inch / 2, topMargin=inch / 2,
                      bottomMargin=inch / 2, title=title).build([t])


def xlsx_sheets(src):
    from openpyxl import load_workbook
    wb = load_workbook(str(src), read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets]
    finally:
        wb.close()


def _sheet_csv_name(src, sheet, count):
    if count == 1:
        return src.stem + ".csv"
    safe = "".join(ch for ch in sheet if ch not in '<>:"/\\|?*')
    return f"{src.stem} - {safe}.csv"


def xlsx_to_csv(src, out_dir):
    sheets = xlsx_sheets(src)
    outs = []
    for name, rows in sheets:
        out = unique_path(out_dir / _sheet_csv_name(src, name, len(sheets)))
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows([["" if v is None else v for v in r] for r in rows])
        outs.append(out)
    return outs


# ---------------------------------------------------------------- Office (COM)
WORD_FMT = {"PDF": 17, "DOCX": 16, "RTF": 6, "ODT": 23, "TXT": 7, "HTML": 10}
EXCEL_FMT = {"XLSX": 51, "ODS": 60, "CSV": 62}
PPT_FMT = {"PDF": 32, "PPTX": 24, "ODP": 35}
_NO_PASSWORD = "\0001drpdf"  # makes protected files fail instead of prompting


class Office:
    """Lazily started, invisible Office apps shared across one batch."""

    def __init__(self):
        self._apps = {}

    def get(self, name):
        if name not in self._apps:
            import win32com.client
            app = win32com.client.DispatchEx(f"{name}.Application")
            try:
                app.DisplayAlerts = 0 if name != "Excel" else False
            except Exception:
                pass
            if name in ("Word", "Excel"):
                try:
                    app.Visible = False
                except Exception:
                    pass
            self._apps[name] = app
        return self._apps[name]

    def close(self):
        for app in self._apps.values():
            try:
                app.Quit()
            except Exception:
                pass
        self._apps.clear()


def word_convert(office, src, label, out):
    word = office.get("Word")
    doc = word.Documents.Open(str(src), ConfirmConversions=False, ReadOnly=True,
                              AddToRecentFiles=False, PasswordDocument=_NO_PASSWORD,
                              Visible=False)
    try:
        if label == "TXT":
            doc.SaveAs2(str(out), FileFormat=WORD_FMT["TXT"], Encoding=65001)
        else:
            doc.SaveAs2(str(out), FileFormat=WORD_FMT[label])
    finally:
        doc.Close(0)


def excel_convert(office, src, label, out_dir):
    xl = office.get("Excel")
    wb = xl.Workbooks.Open(str(src), UpdateLinks=0, ReadOnly=True,
                           Password=_NO_PASSWORD, AddToMru=False)
    outs = []
    try:
        if label == "PDF":
            out = unique_path(out_dir / (src.stem + ".pdf"))
            wb.ExportAsFixedFormat(0, str(out))
            outs.append(out)
        elif label == "CSV":
            count = wb.Worksheets.Count
            for i in range(1, count + 1):
                ws = wb.Worksheets(i)
                out = unique_path(out_dir / _sheet_csv_name(src, ws.Name, count))
                ws.Copy()  # copy the sheet into a new, single-sheet workbook
                tmp = xl.ActiveWorkbook
                try:
                    tmp.SaveAs(str(out), FileFormat=EXCEL_FMT["CSV"])
                finally:
                    tmp.Close(False)
                outs.append(out)
        else:
            out = unique_path(out_dir / (src.stem + TARGETS[label][0]))
            wb.SaveAs(str(out), FileFormat=EXCEL_FMT[label])
            outs.append(out)
    finally:
        wb.Close(False)
    return outs


def ppt_convert(office, src, label, out):
    ppt = office.get("PowerPoint")
    pres = ppt.Presentations.Open(str(src), True, False, False)  # read-only, no window
    try:
        pres.SaveAs(str(out), PPT_FMT[label])
    finally:
        pres.Close()


# ---------------------------------------------------------------- LibreOffice
_LO_FILTER = {"TXT": "txt:Text (encoded):UTF8", "CSV": "csv:Text - txt - csv (StarCalc):44,34,76"}


def soffice_convert(src, label, out):
    tmp = Path(tempfile.mkdtemp())
    try:
        target = _LO_FILTER.get(label, TARGETS[label][0][1:])
        subprocess.run([SOFFICE, "--headless", "--norestore", "--convert-to", target,
                        "--outdir", str(tmp), str(src)], check=True, timeout=300,
                       capture_output=True,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        made = list(tmp.iterdir())
        if not made:
            raise RuntimeError("LibreOffice could not convert this file")
        shutil.move(str(made[0]), str(out))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- dispatcher
def convert(src: Path, label: str, out_dir: Path, opts: dict, office: Office):
    """Convert one file. Returns the list of files that were written."""
    out_dir.mkdir(parents=True, exist_ok=True)
    kind = kind_of(src.suffix)
    ext = TARGETS[label][0]
    out = unique_path(out_dir / (src.stem + ext))

    if kind == "image":
        return image_convert(src, label, out_dir, opts)

    if kind == "pdf":
        if label in ("DOCX", "XLSX", "TXT"):
            src = _repair_if_damaged(src)
            src = _ocr_if_scanned(src, out_dir)
        if label == "DOCX":
            pdf_to_docx(src, out)
        elif label == "XLSX":
            pdf_to_xlsx(src, out)
        elif label == "TXT":
            pdf_to_text(src, out)
        elif label in IMAGE_FMT:
            return pdf_to_images(src, label, out_dir, opts)
        elif label == "PDF":
            shutil.copy2(src, out)
        return [out]

    if kind == "text":
        (text_to_pdf if label == "PDF" else text_to_docx)(src, out)
        return [out]

    if kind == "csv":
        if label == "XLSX":
            csv_to_xlsx(src, out)
        elif label == "CSV":
            csv_to_csv(src, out)
        elif HAS_EXCEL:
            return excel_convert(office, src, label, out_dir)
        else:
            rows_to_pdf(_read_csv_rows(src), out, src.stem)
        return [out]

    if kind == "word":
        if HAS_WORD:
            word_convert(office, src, label, out)
        elif SOFFICE:
            soffice_convert(src, label, out)
        else:
            docx_to_text(src, out)
        return [out]

    if kind == "sheet":
        if HAS_EXCEL:
            return excel_convert(office, src, label, out_dir)
        if SOFFICE and (label != "CSV" or src.suffix.lower() not in (".xlsx", ".xlsm")):
            soffice_convert(src, label, out)
            return [out]
        if label == "CSV":
            return xlsx_to_csv(src, out_dir)
        rows = [r for _, rs in xlsx_sheets(src) for r in rs + [[]]]
        rows_to_pdf(rows, out, src.stem)
        return [out]

    if kind == "slides":
        if label in IMAGE_FMT:  # slides -> pictures via a temporary PDF
            tmp = Path(tempfile.mkdtemp())
            try:
                pdf = tmp / (src.stem + ".pdf")
                if HAS_PPT:
                    ppt_convert(office, src, "PDF", pdf)
                else:
                    soffice_convert(src, "PDF", pdf)
                return pdf_to_images(pdf, label, out_dir, opts)
            finally:
                shutil.rmtree(tmp, ignore_errors=True)
        if HAS_PPT:
            ppt_convert(office, src, label, out)
        else:
            soffice_convert(src, label, out)
        return [out]

    raise RuntimeError("Unsupported file type")


def _repair_if_damaged(src):
    """PDFs whose letters don't say which letters they are would convert to gibberish:
    repair a copy first (see textrepair)."""
    import pymupdf as fitz
    try:
        import textrepair
        with fitz.open(str(src)) as doc:
            if doc.needs_pass or not textrepair.broken_pages(doc):
                return src
            textrepair.repair(doc)
            textrepair.settle_unsure(doc)
            tmp = Path(tempfile.mkdtemp()) / src.name
            doc.save(str(tmp), garbage=1, deflate=True)
    except Exception:
        return src
    _TEMP_DIRS.append(tmp.parent)
    return tmp


def _ocr_if_scanned(src, out_dir):
    """Scanned PDFs have no text to convert; read it with Windows OCR first."""
    import pymupdf as fitz
    try:
        import ocr
    except Exception:
        return src
    with fitz.open(str(src)) as doc:
        if not (ocr.available() and ocr.doc_needs_ocr(doc, max_pages=doc.page_count)):
            return src
        tmp = Path(tempfile.mkdtemp()) / src.name
        with ocr.text_copy(doc) as copy:
            copy.save(str(tmp))
    _TEMP_DIRS.append(tmp.parent)
    return tmp


_TEMP_DIRS = []


def _cleanup_temp():
    for d in _TEMP_DIRS:
        shutil.rmtree(d, ignore_errors=True)


import atexit  # noqa: E402
atexit.register(_cleanup_temp)


def merge_pdfs(pdfs, out):
    import pymupdf as fitz
    merged = fitz.open()
    try:
        for p in pdfs:
            with fitz.open(str(p)) as d:
                merged.insert_pdf(d)
        merged.save(str(out), garbage=3, deflate=True)
    finally:
        merged.close()
